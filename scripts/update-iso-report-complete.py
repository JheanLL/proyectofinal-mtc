import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top='E2E8F0', bottom='E2E8F0', left=None, right=None):
    tcPr = cell._tc.get_or_add_tcPr()
    borders_elm = f'<w:tcBorders {nsdecls("w")}>'
    borders_elm += f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{top}"/>'
    borders_elm += f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{bottom}"/>'
    if left:
        borders_elm += f'<w:left w:val="single" w:sz="4" w:space="0" w:color="{left}"/>'
    else:
        borders_elm += f'<w:left w:val="none"/>'
    if right:
        borders_elm += f'<w:right w:val="single" w:sz="4" w:space="0" w:color="{right}"/>'
    else:
        borders_elm += f'<w:right w:val="none"/>'
    borders_elm += '</w:tcBorders>'
    tcBorders = parse_xml(borders_elm)
    tcPr.append(tcBorders)

def format_cell(cell, text, bold=False, color_rgb=(15, 23, 42), font_size=8.5, fill_hex=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    if fill_hex:
        set_cell_background(cell, fill_hex)
    set_cell_margins(cell)
    set_cell_borders(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Calibri'
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color_rgb)

def main():
    doc_path = '/home/jhean/Desktop/PROYECTO_FINAL_ISO.docx'
    doc = docx.Document(doc_path)

    # Identificar tablas dinámicamente
    t_rf = None
    t_cu = None
    t_bd = None
    t_aud = None
    t_pruebas = None

    for t in doc.tables:
        f0 = t.rows[0].cells[0].text.strip()
        f1 = t.rows[0].cells[1].text.strip() if len(t.rows[0].cells) > 1 else ""

        if 'Requerimiento Funcional' in f1:
            t_rf = t
        elif 'Caso de Uso' in f1:
            t_cu = t
        elif 'Tabla (tbl_)' in f0 and 'Prefijo' in f1:
            t_bd = t
        elif 'Rol de Acceso' in f0 and 'Ámbito de Auditoría' in f1:
            t_aud = t
        elif f0 == 'ID' and 'Módulo Evaluado' in f1:
            t_pruebas = t

    print("Found ISO tables:")
    print("  RF:", t_rf is not None)
    print("  CU:", t_cu is not None)
    print("  BD:", t_bd is not None)
    print("  Auditoria:", t_aud is not None)
    print("  Pruebas:", t_pruebas is not None)

    # 1. Actualizar t_rf: agregar RF-16 y RF-17
    if t_rf:
        existing = [r.cells[0].text.strip() for r in t_rf.rows]
        new_rfs = [
            ('RF-16', 'Filtros de Exclusión en Servidor (tbl_filtro_exclusion)', 
             'Baja administrativa y ocultamiento inmediato (< 0.1 ms) de zonas y frecuencias en memoria de servidor sin alterar APIs externas.', 'Alta (Must - Implementado)'),
            ('RF-17', 'Restauración de Registros desde Auditoría', 
             'Recuperación íntegra de entidades eliminadas desde /admin/auditoria, restableciendo imágenes WebP, coordenadas y tarifas sin pérdida de datos.', 'Alta (Must - Implementado)')
        ]
        for code, name, desc, state in new_rfs:
            if code not in existing:
                row = t_rf.add_row()
                bg_hex = 'FFFFFF' if len(t_rf.rows) % 2 == 0 else 'F8FAFC'
                format_cell(row.cells[0], code, bold=True, fill_hex=bg_hex)
                format_cell(row.cells[1], name, bold=True, fill_hex=bg_hex)
                format_cell(row.cells[2], desc, fill_hex=bg_hex)
                format_cell(row.cells[3], state, bold=True, fill_hex=bg_hex, color_rgb=(16, 185, 129))
                print(f"Added {code} to ISO Table 1.")

    # 2. Actualizar t_cu: agregar Casos de Uso de Filtrado y Restauración
    if t_cu:
        existing_cu = [r.cells[1].text.strip() for r in t_cu.rows]
        new_cus = [
            ('Operador Travel Group / PeruRail / MTC', 'Restaurar Registros desde Bitácora de Auditoría', 
             'Recuperación con un solo clic de zonas o frecuencias eliminadas desde /admin/auditoria con preservación total de fotos WebP y tarifas.'),
            ('Operadores Autorizados', 'Gestión de Filtros de Exclusión en Servidor (tbl_filtro_exclusion)', 
             'Baja lógica inmediata con filtrado en memoria del servidor (< 0.1 ms) para no alterar APIs externas de origen.')
        ]
        for actor, cu_name, cu_desc in new_cus:
            if cu_name not in existing_cu:
                row = t_cu.add_row()
                bg_hex = 'FFFFFF' if len(t_cu.rows) % 2 == 0 else 'F8FAFC'
                format_cell(row.cells[0], actor, bold=True, fill_hex=bg_hex)
                format_cell(row.cells[1], cu_name, bold=True, fill_hex=bg_hex)
                format_cell(row.cells[2], cu_desc, fill_hex=bg_hex)
                print(f"Added {cu_name} to ISO Table 2.")

    # 3. Actualizar t_bd: agregar tbl_filtro_exclusion
    if t_bd:
        existing_tbl = [r.cells[0].text.strip() for r in t_bd.rows]
        if 'tbl_filtro_exclusion' not in existing_tbl:
            row = t_bd.add_row()
            bg_hex = 'FFFFFF' if len(t_bd.rows) % 2 == 0 else 'F8FAFC'
            format_cell(row.cells[0], 'tbl_filtro_exclusion', bold=True, fill_hex=bg_hex)
            format_cell(row.cells[1], 'fil_', fill_hex=bg_hex)
            format_cell(row.cells[2], 'Filtros / Exclusiones', fill_hex=bg_hex)
            format_cell(row.cells[3], 'fil_id, fil_modulo (ZONAS/HORARIOS), fil_registro_id, fil_motivo, fil_usuario_email, fil_activo, fil_fecha_registro', fill_hex=bg_hex)
            print("Added tbl_filtro_exclusion to ISO Table 5.")

    # 4. Actualizar t_aud: agregar regla de RESTORE
    if t_aud:
        # Añadimos fila para evento RESTORE si no existe
        existing_aud = [r.cells[0].text.strip() for r in t_aud.rows]
        if 'Acción RESTORE (Restauración)' not in existing_aud:
            row = t_aud.add_row()
            bg_hex = 'F8FAFC'
            format_cell(row.cells[0], 'Acción RESTORE (Restauración)', bold=True, fill_hex=bg_hex)
            format_cell(row.cells[1], 'Permitido según Rol (RBAC)', fill_hex=bg_hex)
            format_cell(row.cells[2], 'Travel Group: Zonas; PeruRail: Horarios; MTC: Global', fill_hex=bg_hex)
            format_cell(row.cells[3], "POST /api/auditoria/restore reinserta datos, retira de tbl_filtro_exclusion y asienta aud_accion='RESTORE'", fill_hex=bg_hex)
            print("Added RESTORE row to ISO Table 8.")

    # 5. Actualizar t_pruebas: agregar CP-14, CP-15, CP-16
    if t_pruebas:
        existing_cps = [r.cells[0].text.strip() for r in t_pruebas.rows]
        new_cps = [
            ('CP-14', 'Filtro de Exclusión', 'Dar de baja una zona o tren en el servidor.', 'Se registra en tbl_filtro_exclusion y desaparece al instante de la UI y cálculos sin tocar APIs externas. [APROBADO]'),
            ('CP-15', 'Auditoría y Restauración', 'Restaurar registro eliminado desde consola /admin/auditoria.', 'Recupera la entidad con total fidelidad (fotos WebP, coordenadas, tarifas), retira exclusión y genera evento RESTORE. [APROBADO]'),
            ('CP-16', 'Seguridad y RBAC', 'Verificar restricción de permisos por rol en endpoints y restauración.', 'Travel Group restringido a Zonas; PeruRail a Horarios; MTC Admin con acceso global (26/26 pruebas unitarias y de integración aprobadas). [APROBADO]')
        ]
        for cp_id, mod, esc, res in new_cps:
            if cp_id not in existing_cps:
                row = t_pruebas.add_row()
                bg_hex = 'FFFFFF' if len(t_pruebas.rows) % 2 == 0 else 'F8FAFC'
                format_cell(row.cells[0], cp_id, bold=True, fill_hex=bg_hex)
                format_cell(row.cells[1], mod, bold=True, fill_hex=bg_hex)
                format_cell(row.cells[2], esc, fill_hex=bg_hex)
                format_cell(row.cells[3], res, fill_hex=bg_hex, color_rgb=(16, 185, 129))
                print(f"Added {cp_id} to ISO Table 9.")

    # 6. Agregar Historias de Usuario HU-11 y HU-12 en los párrafos de auditoría
    for p in doc.paragraphs:
        if '• HU-10: Historial de Modificaciones Propias' in p.text:
            # Insertar HU-11 y HU-12 después de este párrafo
            doc_text = " ".join([par.text for par in doc.paragraphs])
            if 'HU-11: Restauración de Entidades desde Bitácora de Auditoría' not in doc_text:
                p_hu11 = doc.add_paragraph()
                p._p.addnext(p_hu11._p)
                p_hu11.paragraph_format.space_before = Pt(8)
                p_hu11.paragraph_format.space_after = Pt(4)
                p_hu11.paragraph_format.line_spacing = 1.15
                r11 = p_hu11.add_run(
                    "• HU-11: Restauración de Entidades desde Bitácora de Auditoría (Operadores Travel Group, PeruRail y MTC Admin)\n"
                    "  - Como: Operador institucional autenticado (Travel Group para Zonas, PeruRail para Horarios, MTC para todo).\n"
                    "  - Quiero: Disponer de un botón de restauración interactiva en la bitácora de auditoría para revertir eliminaciones de registros complejos.\n"
                    "  - Para: Recuperar instantáneamente fotografías WebP, coordenadas cartográficas y estructuras tarifarias sin reingresar datos manualmente ni incurrir en pérdidas de información.\n"
                    "  - Criterio de Aceptación: Reversión de la baja, desactivación/retiro de tbl_filtro_exclusion, inserción de evento 'RESTORE' y visualización inmediata en el catálogo público con validación RBAC estricta (HTTP 403 para accesos no autorizados)."
                )
                r11.font.name = 'Calibri'
                r11.font.size = Pt(9)
                r11.font.color.rgb = RGBColor(51, 65, 85)

                p_hu12 = doc.add_paragraph()
                p_hu11._p.addnext(p_hu12._p)
                p_hu12.paragraph_format.space_before = Pt(8)
                p_hu12.paragraph_format.space_after = Pt(4)
                p_hu12.paragraph_format.line_spacing = 1.15
                r12 = p_hu12.add_run(
                    "• HU-12: Filtrado de Exclusión en Servidor a Ultra-Baja Latencia (MTC y Operadores)\n"
                    "  - Como: Administrador del sistema.\n"
                    "  - Quiero: Que las eliminaciones administrativas se procesen mediante la tabla tbl_filtro_exclusion en memoria de servidor (< 0.1 ms).\n"
                    "  - Para: Ocultar registros al instante del público sin modificar destructivamente las APIs externas ni degradar el rendimiento del servidor.\n"
                    "  - Criterio de Aceptación: Filtrado server-side en menos de 0.1 ms que excluye registros activos de las consultas públicas y asesor turístico."
                )
                r12.font.name = 'Calibri'
                r12.font.size = Pt(9)
                r12.font.color.rgb = RGBColor(51, 65, 85)
                print("Inserted HU-11 and HU-12 in ISO document.")
            break

    doc.save(doc_path)
    print("PROYECTO_FINAL_ISO.docx updated and saved successfully!")

if __name__ == '__main__':
    main()
