import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top='E2E8F0', bottom='E2E8F0'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders_elm = f'<w:tcBorders {nsdecls("w")}>'
    borders_elm += f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{top}"/>'
    borders_elm += f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{bottom}"/>'
    borders_elm += f'<w:left w:val="none"/>'
    borders_elm += f'<w:right w:val="none"/>'
    borders_elm += '</w:tcBorders>'
    tcBorders = parse_xml(borders_elm)
    tcPr.append(tcBorders)

def format_cell(cell, text, font_size=8, fill_hex=None):
    if fill_hex:
        set_cell_background(cell, fill_hex)
    set_cell_margins(cell)
    set_cell_borders(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.font.name = 'Consolas'
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(15, 23, 42)

NEW_STATIONS_SQL = """-- ------------------------------------------------------------------------------
-- 2. SEED OFICIAL: CATÁLOGO MAESTRO DE ESTACIONES FERROVIARIAS (PeruRail / MTC)
-- Nota Arquitectónica: Este seed garantiza la Integridad Referencial (Foreign Keys)
-- en MySQL (tbl_zona_turistica y tbl_horario_tren). En tiempo de ejecución, el cliente
-- web consume esta infraestructura dinámicamente mediante el endpoint REST GET /api/estaciones.
-- ------------------------------------------------------------------------------
INSERT INTO tbl_estacion (est_id, est_codigo, est_nombre, est_ciudad, est_departamento, est_altitud_msnm, est_latitud, est_longitud, est_descripcion, est_servicios, est_imagen_url)
VALUES
('est_01', 'EST-CUS-SP', 'Estación San Pedro (Cusco)', 'Cusco', 'Cusco', 3399, -13.52040000, -71.98280000, 'Ubicada en el corazón histórico de Cusco, frente al célebre Mercado Central y a escasas cuadras de la Plaza de Armas.', '["Boletería PeruRail", "Sala de espera VIP", "Información turística iPerú", "Custodia de equipaje", "Cafetería andina", "Wifi gratuito"]', 'https://images.unsplash.com/photo-1589553416260-f586c8f1514f?auto=format&fit=crop&w=1200&q=80'),
('est_02', 'EST-CUS-POR', 'Estación Poroy (Cusco)', 'Poroy', 'Cusco', 3499, -13.49120000, -72.04650000, 'Punto de partida ideal en las afueras de la ciudad imperial de Cusco hacia el Valle Sagrado y Machu Picchu.', '["Boletería", "Estacionamiento vigilado", "Servicio de maleteros", "Cafetería", "Servicios higiénicos accesibles"]', 'https://images.unsplash.com/photo-1526392060635-9d6019884377?auto=format&fit=crop&w=1200&q=80'),
('est_03', 'EST-VAL-OLL', 'Estación Ollantaytambo', 'Ollantaytambo', 'Cusco', 2792, -13.25890000, -72.26380000, 'Principal nodo ferroviario del Valle Sagrado con conexión directa diaria hacia Machu Picchu.', '["Boletería automatizada", "Salas de embarque", "Restaurantes", "Módulo SENAMHI", "Artesanías locales", "Parqueo de bicicletas"]', 'https://images.unsplash.com/photo-1587595431973-160d0d94add1?auto=format&fit=crop&w=1200&q=80'),
('est_04', 'EST-MAC-AGU', 'Estación Machu Picchu Pueblo (Aguas Calientes)', 'Machu Picchu Pueblo', 'Cusco', 2040, -13.15490000, -72.52550000, 'Estación terminal a orillas del río Vilcanota, puerta de entrada peatonal y vehicular al Santuario Histórico de Machu Picchu.', '["Centro de atención al visitante", "Boletería", "Embarque preferencial", "Asistencia médica", "Oficina MTC / Dircetur"]', 'https://images.unsplash.com/photo-1509299349698-dd22323b5963?auto=format&fit=crop&w=1200&q=80'),
('est_05', 'EST-VAL-URU', 'Estación Urubamba', 'Urubamba', 'Cusco', 2870, -13.30560000, -72.11580000, 'Estación rodeada de clima templado y valles fértiles, conectada con hoteles boutique y huertos orgánicos.', '["Boletería", "Salón de espera", "Jardín andino", "Información turística", "Café bar"]', 'https://images.unsplash.com/photo-1609137144813-7d9921338f24?auto=format&fit=crop&w=1200&q=80'),
('est_06', 'EST-PUN-TIT', 'Estación Puno (Lago Titicaca)', 'Puno', 'Puno', 3827, -15.83640000, -70.02190000, 'Estación histórica a orillas del lago navegable más alto del mundo, conectando el Altiplano con Cusco.', '["Boletería Titicaca Train", "Sala de embarque folclórica", "Oxígeno de cortesía", "Custodia de equipaje"]', 'https://images.unsplash.com/photo-1534008897995-27a23e859048?auto=format&fit=crop&w=1200&q=80')
ON DUPLICATE KEY UPDATE est_nombre = VALUES(est_nombre);"""

def main():
    # 1. Actualizar PROYECTO_FINAL_ISO.docx
    doc_path = '/home/jhean/Desktop/PROYECTO_FINAL_ISO.docx'
    doc = docx.Document(doc_path)

    for t in doc.tables:
        f0 = t.rows[0].cells[0].text.strip()
        if 'Script DDL Maestro' in f0:
            ddl_text = t.rows[0].cells[0].text
            # Reemplazar bloque antiguo de estaciones
            old_start = "INSERT INTO tbl_estacion (est_id, est_codigo"
            old_end = "ON DUPLICATE KEY UPDATE est_nombre = VALUES(est_nombre);"
            if old_start in ddl_text and old_end in ddl_text:
                idx1 = ddl_text.find(old_start)
                idx2 = ddl_text.find(old_end, idx1) + len(old_end)
                ddl_text = ddl_text[:idx1] + NEW_STATIONS_SQL + ddl_text[idx2:]
                format_cell(t.rows[0].cells[0], ddl_text, fill_hex='F8FAFC')
                print("Replaced stations in PROYECTO_FINAL_ISO.docx Table 7.")
                break

    doc.save(doc_path)
    print("Saved:", doc_path)

    # 2. Actualizar database_schema_aiven.sql
    sql_path = '/home/jhean/Desktop/Only Software/proyectoSoftware/proyectofinal/database_schema_aiven.sql'
    with open(sql_path, 'r', encoding='utf-8') as f:
        sql_content = f.read()

    if "INSERT INTO tbl_estacion (est_id, est_codigo" in sql_content:
        idx1 = sql_content.find("INSERT INTO tbl_estacion (est_id, est_codigo")
        idx2 = sql_content.find("ON DUPLICATE KEY UPDATE est_nombre = VALUES(est_nombre);", idx1) + len("ON DUPLICATE KEY UPDATE est_nombre = VALUES(est_nombre);")
        sql_content = sql_content[:idx1] + NEW_STATIONS_SQL + sql_content[idx2:]
        with open(sql_path, 'w', encoding='utf-8') as f:
            f.write(sql_content)
        print("Updated database_schema_aiven.sql with exact 6 real stations.")

if __name__ == '__main__':
    main()
