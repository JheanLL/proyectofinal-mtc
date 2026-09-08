import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top='E2E8F0', bottom='E2E8F0'):
    tcPr = cell._tc.get_or_add_tcPr()
    borders_elm = f'<w:tcBorders {nsdecls("w")}>'
    borders_elm += f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{top}"/>'
    borders_elm += f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{bottom}"/>'
    borders_elm += f'<w:left w:val="none"/>'
    borders_elm += f'<w:right w:val="none"/>'
    borders_elm += '</w:tcBorders>'
    tcBorders = parse_xml(borders_elm)
    tcPr.append(tcBorders)

def format_cell(cell, text, bold=False, color_rgb=(15, 23, 42), font_size=8.5, fill_hex=None, align=WD_ALIGN_PARAGRAPH.LEFT):
    if fill_hex:
        set_cell_background(cell, fill_hex)
    set_cell_margins(cell)
    set_cell_borders(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Calibri'
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color_rgb)

API_CATALOG = [
    ("PeruRail / MTC (Estaciones)", "GET", 
     "/api/estaciones\nhttps://proyectofinal-mtc.vercel.app/api/estaciones",
     "Ninguno (Opcional: ?id=est_01)",
     "Catálogo maestro de las 6 estaciones oficiales con coordenadas GPS, altitud msnm y servicios.",
     "Público (Libre)"),
    ("SENAMHI (Clima en Vivo)", "GET", 
     "/api/senamhi\nhttps://proyectofinal-mtc.vercel.app/api/senamhi",
     "?estacion_id=est_01\n?refresh=true (Bypass)",
     "Previsión meteorológica satelital en tiempo real (Open-Meteo): temperatura, humedad, lluvia %, índice UV y alertas. Caché 5 min.",
     "Público (Libre)"),
    ("PeruRail (Horarios y Trenes)", "GET, POST, PUT, DELETE", 
     "/api/horarios\nhttps://proyectofinal-mtc.vercel.app/api/horarios",
     "?origen=est_01&destino=est_04\nPayload JSON (CRUD)",
     "Frecuencias, horarios de salida/llegada, tipos de servicio (Expedition, Vistadome) y tarifas en PEN y USD.",
     "GET: Público\nCRUD: PERURAIL / ADMIN"),
    ("Travel Group (Zonas Turísticas)", "GET, POST, PUT, DELETE", 
     "/api/zonas\nhttps://proyectofinal-mtc.vercel.app/api/zonas",
     "?estacion_id=est_01&categoria=naturaleza\nPayload JSON con imagen WebP",
     "Circuitos peatonales ida/vuelta, tiempos, distancias, dificultad, georreferencias y fotos WebP. Filtra exclusiones.",
     "GET: Público\nCRUD: TRAVEL_GROUP / ADMIN"),
    ("MTC (Sincronización de APIs)", "GET, POST", 
     "/api/sync\nhttps://proyectofinal-mtc.vercel.app/api/sync",
     "POST: Token JWT en header/cookie",
     "Monitoreo de latencia (ms) y estado de salud de los 3 conectores. El POST fuerza actualización inmediata de flujos e invalida caché.",
     "GET: Público\nPOST: Solo ADMIN MTC"),
    ("MTC / Seguridad (Auditoría)", "GET", 
     "/api/auditoria\nhttps://proyectofinal-mtc.vercel.app/api/auditoria",
     "?modulo=ZONAS&accion=DELETE\nToken JWT",
     "Bitácora inmutable no repudiable con fecha, usuario, IP, módulo, acción y snapshot JSON del registro afectado.",
     "ADMIN: Global (100%)\nOperadores: Propio"),
    ("MTC / Operadores (Restauración)", "POST", 
     "/api/auditoria/restore\nhttps://proyectofinal-mtc.vercel.app/api/auditoria/restore",
     "Body: { aud_id: 123 }\nToken JWT",
     "Restaura entidades eliminadas desde la bitácora recuperando fotos WebP, coordenadas y tarifas; des-excluye de tbl_filtro_exclusion.",
     "RBAC según módulo:\nTRAVEL_GROUP / PERURAIL / ADMIN"),
    ("MTC / Seguridad (Login RBAC)", "POST, GET", 
     "/api/auth/login | /api/auth/me\nhttps://proyectofinal-mtc.vercel.app/api/auth/login",
     "Body: { email, password }",
     "Autenticación con bcrypt y emisión de cookie JWT HttpOnly segura para los 3 perfiles oficiales.",
     "Público (Login)"),
    ("MTC / Turistas (Itinerarios)", "GET, POST", 
     "/api/itinerarios\nhttps://proyectofinal-mtc.vercel.app/api/itinerarios",
     "?codigo=MTC-xxxx\nBody: { resumenItinerario }",
     "Almacenamiento y recuperación anónima de itinerarios planificados mediante token alfanumérico único para compartir o imprimir.",
     "Público (Libre)"),
    ("MTC / Infraestructura (Healthcheck)", "GET", 
     "/api/db-status\nhttps://proyectofinal-mtc.vercel.app/api/db-status",
     "Ninguno",
     "Verificación de conectividad SSL en tiempo real con el clúster gestionado de Aiven for MySQL y estado de las 10 tablas.",
     "Público (Diagnóstico)")
]

def add_api_catalog_to_doc(doc_path):
    doc = docx.Document(doc_path)
    doc_text = " ".join([p.text for p in doc.paragraphs])
    
    if "Catálogo Oficial de APIs REST y Endpoints en Producción" in doc_text:
        print(f"API catalog already in {doc_path}")
        return

    # Buscar la tabla de enlaces oficiales (Table 10) o el párrafo del Módulo 4
    target_element = None
    for t in doc.tables:
        f0 = t.rows[0].cells[0].text.strip()
        if 'ENLACES OFICIALES DEL SOFTWARE' in f0:
            target_element = t._tbl
            break

    if not target_element:
        for p in doc.paragraphs:
            if '4. Módulo de Informes:' in p.text:
                target_element = p._p
                break

    if target_element is not None:
        # Añadir título
        p_title = doc.add_paragraph()
        target_element.addnext(p_title._p)
        p_title.paragraph_format.space_before = Pt(14)
        p_title.paragraph_format.space_after = Pt(6)
        r_t = p_title.add_run("Catálogo Oficial de APIs REST y Endpoints en Producción (Cloud Deployment)")
        r_t.bold = True
        r_t.font.name = 'Calibri'
        r_t.font.size = Pt(10)
        r_t.font.color.rgb = RGBColor(15, 23, 42)

        # Crear tabla de APIs
        api_table = doc.add_table(rows=1, cols=6)
        p_title._p.addnext(api_table._tbl)

        headers = [
            "Flujo / Módulo Oficial", 
            "Método", 
            "Endpoint y URL en Producción", 
            "Parámetros / Entradas", 
            "Descripción y Salida", 
            "Nivel RBAC"
        ]
        col_widths = [Inches(1.3), Inches(0.8), Inches(2.2), Inches(1.4), Inches(2.0), Inches(1.1)]

        for i, h in enumerate(headers):
            format_cell(api_table.rows[0].cells[i], h, bold=True, color_rgb=(255, 255, 255), fill_hex='0F172A', font_size=8)
            api_table.rows[0].cells[i].width = col_widths[i]

        for row_idx, item in enumerate(API_CATALOG):
            row = api_table.add_row()
            bg_hex = 'FFFFFF' if row_idx % 2 == 0 else 'F8FAFC'
            for col_idx, text in enumerate(item):
                cell = row.cells[col_idx]
                cell.width = col_widths[col_idx]
                is_bold = (col_idx in [0, 1])
                color = (16, 185, 129) if text == 'Público (Libre)' else (15, 23, 42)
                format_cell(cell, text, bold=is_bold, fill_hex=bg_hex, color_rgb=color, font_size=7.5)

        # Actualizar descripción del Módulo 1 (Integración) con los links directos
        for p in doc.paragraphs:
            if p.text.strip().startswith('• 1. Módulo de Integración de Datos:'):
                p.text = (
                    "• 1. Módulo de Integración de Datos (APIs REST): Conectores con flujos periódicos oficiales:\n"
                    "  - PeruRail Estaciones: https://proyectofinal-mtc.vercel.app/api/estaciones (Catálogo de 6 estaciones con GPS y altitud)\n"
                    "  - PeruRail Horarios: https://proyectofinal-mtc.vercel.app/api/horarios (Frecuencias ferroviarias y tarifas en PEN y USD)\n"
                    "  - SENAMHI Clima: https://proyectofinal-mtc.vercel.app/api/senamhi (Previsiones en tiempo real vía Open-Meteo satelital)\n"
                    "  - Travel Group Zonas: https://proyectofinal-mtc.vercel.app/api/zonas (Circuitos a pie, georreferencias e imágenes WebP)\n"
                    "  - Sincronizador MTC: https://proyectofinal-mtc.vercel.app/api/sync (Latencia en ms y sincronización periódica)"
                )
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(6)
                p.paragraph_format.line_spacing = 1.15
                break

        doc.save(doc_path)
        print(f"Added complete API directory and links to {doc_path}")

if __name__ == '__main__':
    add_api_catalog_to_doc('/home/jhean/Desktop/PROYECTO_FINAL_ISO.docx')
    add_api_catalog_to_doc('/home/jhean/Desktop/INFORME_PRACTICA_ZONAS_TURISTICAS.docx')
