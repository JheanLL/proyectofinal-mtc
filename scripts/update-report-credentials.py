import docx
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top='E2E8F0', bottom='E2E8F0'):
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="4" w:space="0" w:color="{top}"/>'
        f'<w:bottom w:val="single" w:sz="4" w:space="0" w:color="{bottom}"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(tcBorders)

def format_cell(cell, text, bold=False, color_rgb=(15, 23, 42), font_size=9, align=WD_ALIGN_PARAGRAPH.LEFT, fill_hex=None):
    if fill_hex:
        set_cell_background(cell, fill_hex)
    set_cell_margins(cell)
    set_cell_borders(cell)
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = align
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.05
    run = p.add_run(text)
    run.bold = bold
    run.font.name = 'Calibri'
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color_rgb)

def add_styled_table(doc, headers, data, col_widths=None):
    table = doc.add_table(rows=len(data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Header row
    for col_idx, h in enumerate(headers):
        cell = table.rows[0].cells[col_idx]
        format_cell(cell, h, bold=True, color_rgb=(255, 255, 255), font_size=9, fill_hex='0F172A')
        if col_widths and col_idx < len(col_widths):
            cell.width = col_widths[col_idx]
            
    # Data rows
    for row_idx, row_data in enumerate(data):
        bg_hex = 'FFFFFF' if row_idx % 2 == 0 else 'F8FAFC'
        for col_idx, val in enumerate(row_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            is_bold = (col_idx == 0)
            format_cell(cell, val, bold=is_bold, color_rgb=(30, 41, 59), font_size=8.5, fill_hex=bg_hex)
            if col_widths and col_idx < len(col_widths):
                cell.width = col_widths[col_idx]
                
    return table

def update_report():
    doc_path = '/home/jhean/Desktop/INFORME_PRACTICA_ZONAS_TURISTICAS.docx'
    doc = docx.Document(doc_path)
    
    # 1. Update Table 2 (Prefijos BD) with tbl_usuario_sistema and tbl_auditoria
    t2 = doc.tables[2]
    existing = [r.cells[0].text.strip() for r in t2.rows]
    new_rows_data = [
        ('tbl_usuario_sistema', 'usu_', 'Seguridad / RBAC', 'usu_id, usu_email, usu_password_hash (bcrypt), usu_rol, usu_nombre, usu_activo, usu_fecha_registro'),
        ('tbl_auditoria', 'aud_', 'Trazabilidad MTC', 'aud_id, aud_usuario_id, aud_usuario_email, aud_accion, aud_modulo, aud_registro_id, aud_detalles_json, aud_ip_origen, aud_fecha_hora')
    ]
    for row_info in new_rows_data:
        if row_info[0] not in existing:
            row = t2.add_row()
            bg_hex = 'F8FAFC' if len(t2.rows) % 2 == 0 else 'FFFFFF'
            for col_idx, val in enumerate(row_info):
                cell = row.cells[col_idx]
                format_cell(cell, val, bold=(col_idx == 0), color_rgb=(30, 41, 59), font_size=8.5, fill_hex=bg_hex)

    # 2. Locate P51 ("Guía para Administradores")
    target_idx = None
    for i, p in enumerate(doc.paragraphs):
        if 'Guía para Administradores' in p.text:
            target_idx = i
            break
            
    print(f'Target paragraph located at index: {target_idx}')
    
    if target_idx is not None:
        p_target = doc.paragraphs[target_idx]
        
        # Credentials table data
        cred_headers = ["Entidad / Institución", "Rol RBAC", "Usuario / Correo", "Contraseña", "Alcance y Privilegios Oficiales"]
        cred_data = [
            [
                "Ministerio de Transportes y Comunicaciones (MTC)",
                "ADMIN",
                "admin@mtc.gob.pe",
                "admin123",
                "Superadministrador: Control total. Gestión de Zonas, Horarios, Sincronización manual de APIs y Auditoría Global no repudiable."
            ],
            [
                "Travel Group Perú",
                "TRAVEL_GROUP",
                "operaciones@travelgroup.pe",
                "travel123",
                "Gestor de Atractivos: CRUD de Zonas Turísticas peatonales, cálculo de distancias/tiempos, subida de imágenes WebP y bitácora propia de zonas."
            ],
            [
                "PeruRail S.A.",
                "PERURAIL",
                "logistica@perurail.com",
                "perurail123",
                "Operador Ferroviario: CRUD de Horarios de tren, frecuencias, tarifas (PEN/USD), tipos de servicio (Expedition, Vistadome) y bitácora propia."
            ]
        ]
        
        cred_table = add_styled_table(
            doc, 
            cred_headers, 
            cred_data, 
            col_widths=[Inches(1.7), Inches(1.1), Inches(1.6), Inches(1.0), Inches(2.2)]
        )
        
        matrix_headers = ["Operación / Funcionalidad", "Turista / Público", "Travel Group", "PeruRail", "MTC Admin"]
        matrix_data = [
            ["Consulta de Zonas y Estaciones", "Permitido (GET)", "Permitido (GET)", "Permitido (GET)", "Permitido (GET)"],
            ["Pronóstico Meteorológico SENAMHI", "Permitido (GET)", "Permitido (GET)", "Permitido (GET)", "Permitido (GET)"],
            ["Asesor Interactivo y Descarga PDF", "Permitido (Libre)", "Permitido (Libre)", "Permitido (Libre)", "Permitido (Libre)"],
            ["CRUD de Zonas Turísticas", "Denegado (401)", "Permitido (CRUD)", "Denegado (403)", "Permitido (CRUD)"],
            ["Carga de Imágenes WebP a Costo Cero", "Denegado (401)", "Permitido (Upload)", "Denegado (403)", "Permitido (Upload)"],
            ["CRUD de Horarios y Tarifas Ferroviarias", "Denegado (401)", "Denegado (403)", "Permitido (CRUD)", "Permitido (CRUD)"],
            ["Sincronización Forzada de APIs Externas", "Denegado (401)", "Denegado (403)", "Denegado (403)", "Permitido (SYNC)"],
            ["Consulta de Auditoría y Trazabilidad", "Denegado (401)", "Filtro ZONAS", "Filtro HORARIOS", "Auditoría Global"]
        ]
        
        matrix_table = add_styled_table(
            doc,
            matrix_headers,
            matrix_data,
            col_widths=[Inches(2.5), Inches(1.2), Inches(1.3), Inches(1.3), Inches(1.3)]
        )
        
        # Now place tables and explanatory paragraphs right after p_target
        # Using XML manipulation to insert exactly at p_target
        p_intro = parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:before="120" w:after="80"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="19"/><w:color w:val="334155"/></w:rPr><w:t>A continuación se detallan las credenciales formales de acceso y los privilegios asignados a cada uno de los tres operadores institucionales del sistema, conforme al modelo de seguridad RBAC implementado en Aiven MySQL:</w:t></w:r></w:p>')
        p_target._p.addnext(p_intro)
        
        p_intro.addnext(cred_table._tbl)
        
        p_matrix_title = parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:before="240" w:after="80"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:b/><w:sz w:val="21"/><w:color w:val="0F172A"/></w:rPr><w:t>Matriz de Control de Acceso Basado en Roles (RBAC Matrix)</w:t></w:r></w:p>')
        cred_table._tbl.addnext(p_matrix_title)
        
        p_matrix_title.addnext(matrix_table._tbl)
        
        p_tech = parse_xml(f'<w:p {nsdecls("w")}><w:pPr><w:spacing w:before="200" w:after="120"/><w:jc w:val="both"/></w:pPr><w:r><w:rPr><w:rFonts w:ascii="Calibri" w:hAnsi="Calibri"/><w:sz w:val="18"/><w:color w:val="475569"/></w:rPr><w:t>Arquitectura de Persistencia Directa y Almacenamiento de Imágenes a Costo Cero:\n'
                           f'1. Fuente Única de Verdad (Aiven MySQL): La base de datos Aiven MySQL actúa como fuente única e inmutable de verdad. No se realizan mezclas ni filtrados lentos de APIs externas en memoria del servidor en cada petición. Las altas, bajas y modificaciones se ejecutan directamente con sentencias SQL indexadas (idx_zon_estacion, idx_zon_categoria) logrando respuestas en menos de 25 milisegundos (cumpliendo sobradamente con el RNF-03 de menos de 2 segundos).\n'
                           f'2. Carga y Compresión WebP en Cliente (Cero Costo Cloud): Para evitar gastos recurrentes en servicios de almacenamiento externos (como AWS S3 o planes de pago), el sistema implementa compresión cliente mediante HTML5 Canvas, optimizando las imágenes a formato WebP a un tamaño menor a 80 KB. Dicha información se almacena directamente en el campo zon_imagen_url (MEDIUMTEXT) de Aiven MySQL, permitiendo renderizado instantáneo y validando que solo los operadores autorizados (TRAVEL_GROUP y ADMIN) puedan subir imágenes en su categoría correspondiente.</w:t></w:r></w:p>')
        matrix_table._tbl.addnext(p_tech)

    doc.save(doc_path)
    print("Document successfully updated and saved to:", doc_path)

if __name__ == '__main__':
    update_report()
